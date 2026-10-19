"""Market-cap classification.

Convention (AMFI / SEBI): rank 1-100 by full market cap = large-cap, 101-250 = mid-cap,
251 onward = small-cap. The ranks come from the *current* AMFI list, which AMFI republishes
half-yearly. Import each new list as a new classification version; never edit an old one.
"""

import csv
import logging
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from sqlalchemy import select, update
from sqlalchemy.orm import Session

from app.models import CapCategory, ClassificationVersion, StockUniverse
from app.services.market_data.base import MarketDataProvider

log = logging.getLogger(__name__)

LARGE_CAP_MAX_RANK = 100
MID_CAP_MAX_RANK = 250


def category_for_rank(rank: int) -> CapCategory:
    if rank < 1:
        raise ValueError(f"invalid market-cap rank {rank}")
    if rank <= LARGE_CAP_MAX_RANK:
        return CapCategory.LARGE
    if rank <= MID_CAP_MAX_RANK:
        return CapCategory.MID
    return CapCategory.SMALL


@dataclass(frozen=True)
class ClassifiedStock:
    symbol: str
    company_name: str
    rank: int
    isin: str | None = None
    official_category: CapCategory | None = None  # AMFI's own column, when the file has one

    @property
    def category(self) -> CapCategory:
        return self.official_category or category_for_rank(self.rank)


def _norm(header) -> str:
    return "".join(ch for ch in str(header or "").lower() if ch.isalnum())


_RANK_HEADERS = {"srno", "sno", "rank", "srnumber"}
_NAME_HEADERS = {"companyname", "name", "company"}
_SYMBOL_HEADERS = {"nsesymbol", "symbol", "nsecode"}
_ISIN_HEADERS = {"isin"}
_CATEGORY_PREFIX = "categorization"
_NO_SYMBOL = {"", "-", "NA", "N/A", "NONE"}
_CATEGORY_NAMES = {"largecap": CapCategory.LARGE, "midcap": CapCategory.MID, "smallcap": CapCategory.SMALL}


def _parse_rows(header: list, rows) -> list[ClassifiedStock]:
    cols = {_norm(h): i for i, h in enumerate(header) if h is not None}

    def pick(options: set[str], required: bool) -> int | None:
        for key in options:
            if key in cols:
                return cols[key]
        if required:
            raise ValueError(f"classification file is missing a column like {sorted(options)}")
        return None

    rank_col = pick(_RANK_HEADERS, True)
    name_col = pick(_NAME_HEADERS, True)
    sym_col = pick(_SYMBOL_HEADERS, True)
    isin_col = pick(_ISIN_HEADERS, False)
    cat_col = next((i for k, i in cols.items() if k.startswith(_CATEGORY_PREFIX)), None)

    def cell(row, i):
        return row[i] if i is not None and i < len(row) and row[i] is not None else ""

    out: list[ClassifiedStock] = []
    seen: set[str] = set()
    mismatches = 0
    for row in rows:
        symbol = str(cell(row, sym_col)).strip().upper()
        rank_raw = str(cell(row, rank_col)).strip()
        if symbol in _NO_SYMBOL or not rank_raw:
            continue  # not listed on NSE, or a blank/footer row
        if symbol in seen:
            raise ValueError(f"duplicate NSE symbol {symbol} in classification file")
        seen.add(symbol)
        rank = int(float(rank_raw))
        official = _CATEGORY_NAMES.get(_norm(cell(row, cat_col))) if cat_col is not None else None
        if official and official != category_for_rank(rank):
            mismatches += 1
        out.append(
            ClassifiedStock(
                symbol=symbol,
                company_name=str(cell(row, name_col) or symbol).strip(),
                rank=rank,
                isin=str(cell(row, isin_col)).strip() or None,
                official_category=official,
            )
        )
    if mismatches:
        log.warning("AMFI category differs from rank-based category", extra={"rows": mismatches})
    return out


def parse_amfi_csv(path: str | Path) -> list[ClassifiedStock]:
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    if not rows:
        raise ValueError("classification file has no header row")
    return _parse_rows(rows[0], rows[1:])


def parse_amfi_xlsx(path: str | Path) -> list[ClassifiedStock]:
    """AMFI's official spreadsheet: a title row, then the header row ('Sr. No.', ...)."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    finally:
        wb.close()
    header_at = next((i for i, r in enumerate(rows[:20]) if any(_norm(c) in _RANK_HEADERS for c in r)), None)
    if header_at is None:
        raise ValueError("could not find the header row (expected a 'Sr. No.' column)")
    return _parse_rows(list(rows[header_at]), rows[header_at + 1 :])


def parse_amfi_file(path: str | Path) -> list[ClassifiedStock]:
    """Parse an AMFI market-cap list (.xlsx as published, or .csv). Rows without an NSE symbol are skipped."""
    return parse_amfi_xlsx(path) if str(path).lower().endswith((".xlsx", ".xlsm")) else parse_amfi_csv(path)


def save_classification(
    db: Session,
    stocks: list[ClassifiedStock],
    version: str,
    source: str,
    effective_date: date,
    is_mock: bool = False,
    make_current: bool = True,
) -> ClassificationVersion:
    if db.get(ClassificationVersion, version):
        raise ValueError(f"classification version {version} already exists; versions are immutable")
    if make_current:
        db.execute(update(ClassificationVersion).values(is_current=False))
    cv = ClassificationVersion(
        version=version, source=source, effective_date=effective_date, is_current=make_current, is_mock=is_mock
    )
    db.add(cv)
    db.flush()
    for s in stocks:
        db.add(
            StockUniverse(
                symbol=s.symbol,
                company_name=s.company_name,
                isin=s.isin,
                cap_category=s.category.value,
                cap_rank=s.rank,
                classification_version=version,
                classification_date=effective_date,
                is_active=True,
            )
        )
    db.commit()
    log.info("classification imported", extra={"version": version, "count": len(stocks), "source": source})
    return cv


def import_mock_classification(db: Session, provider: MarketDataProvider, effective_date: date) -> ClassificationVersion:
    """Rank the mock provider's fictional universe in listing order. Development only."""
    if not provider.capabilities.is_mock:
        raise ValueError("mock classification may only be used with the mock provider")
    stocks = [
        ClassifiedStock(symbol=u.symbol, company_name=u.company_name, rank=i)
        for i, u in enumerate(provider.list_universe(), start=1)
    ]
    return save_classification(
        db, stocks, version=f"mock-{effective_date.isoformat()}", source="mock (synthetic ranks)",
        effective_date=effective_date, is_mock=True,
    )


def current_universe(db: Session) -> tuple[ClassificationVersion | None, list[StockUniverse]]:
    cv = db.scalar(select(ClassificationVersion).where(ClassificationVersion.is_current.is_(True)))
    if cv is None:
        return None, []
    rows = db.scalars(
        select(StockUniverse)
        .where(StockUniverse.classification_version == cv.version, StockUniverse.is_active.is_(True))
        .order_by(StockUniverse.symbol)
    ).all()
    return cv, list(rows)
